"""Compile full dissertation (chapters 1-9) into a single Word document."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parent.parent
DISS_DIR = ROOT / "dissertation"
OUTPUT = ROOT / "docs" / "submission" / "NeuroScan_Nepal_Dissertation_Full.docx"

CHAPTERS = [
    ("00_Front_Matter.md", "Front Matter"),
    ("Chapter_01_Introduction.md", "Chapter 1"),
    ("Chapter_02_Literature_Review.md", "Chapter 2"),
    ("Chapter_03_Requirements.md", "Chapter 3"),
    ("Chapter_04_Design.md", "Chapter 4"),
    ("Chapter_05_Implementation.md", "Chapter 5"),
    ("Chapter_06_Testing.md", "Chapter 6"),
    ("Chapter_07_Evaluation.md", "Chapter 7"),
    ("Chapter_08_Ethics.md", "Chapter 8"),
    ("Chapter_09_Conclusion.md", "Chapter 9"),
]

STUDENT = "Pragya Gajurel"
STUDENT_ID = "23189634"


def _add_title_page(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("NeuroScan Nepal\n")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor(0x0D, 0x94, 0x88)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run(
        "An AI-Assisted Brain MRI Screening Platform for Nepal\n\n"
        "Final Year Project Dissertation\n"
        "Full Draft (Chapters 1–9)\n"
    )
    sub_run.font.size = Pt(14)

    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    info.add_run(
        f"\n{STUDENT} ({STUDENT_ID})\n"
        "BSc (Hons) Computer and Data Science\n"
        "Birmingham City University · Sunway College Kathmandu\n"
    )
    doc.add_page_break()


def _parse_markdown_line(doc: Document, line: str, in_code: bool) -> bool:
    stripped = line.rstrip()

    if stripped.startswith("```"):
        return not in_code

    if in_code:
        p = doc.add_paragraph(stripped)
        p.style = "Intense Quote"
        return True

    if not stripped:
        return False

    if stripped.startswith("# "):
        doc.add_heading(stripped[2:], level=1)
    elif stripped.startswith("## "):
        doc.add_heading(stripped[3:], level=2)
    elif stripped.startswith("### "):
        doc.add_heading(stripped[4:], level=3)
    elif stripped.startswith("---"):
        doc.add_paragraph("─" * 40)
    elif stripped.startswith("|") and stripped.endswith("|"):
        doc.add_paragraph(stripped, style="List Bullet")
    elif stripped.startswith("- ") or stripped.startswith("* "):
        doc.add_paragraph(stripped[2:], style="List Bullet")
    elif stripped.startswith("> "):
        p = doc.add_paragraph(stripped[2:])
        p.paragraph_format.left_indent = Inches(0.5)
    else:
        doc.add_paragraph(stripped)

    return False


def _render_chapter(doc: Document, md_path: Path) -> None:
    if not md_path.exists():
        doc.add_paragraph(f"[Missing: {md_path.name}]")
        return

    text = md_path.read_text(encoding="utf-8")
    in_code = False
    for line in text.splitlines():
        in_code = _parse_markdown_line(doc, line, in_code)

    doc.add_page_break()


def main() -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    _add_title_page(doc)

    for filename, _label in CHAPTERS:
        _render_chapter(doc, DISS_DIR / filename)

    doc.save(OUTPUT)
    print(f"Saved: {OUTPUT}")


if __name__ == "__main__":
    main()
