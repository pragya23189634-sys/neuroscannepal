"""Generate NeuroScan Nepal 32-week logbook Word document."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path(__file__).resolve().parent.parent / "docs" / "submission" / "NeuroScan_Nepal_32_Week_Logbook_Updated.docx"

WEEKS = [
    (1, "Project Initiation & Problem Definition", "Phase 1: Research & Planning",
     ["Research brain tumour screening challenges in Nepal",
      "Review existing MRI AI systems and literature",
      "Define problem statement and project scope",
      "Draft project title: NeuroScan Nepal",
      "Discuss Grande International Hospital dataset as primary data source",
      "Meet supervisor and agree timeline"],
     ["Problem statement (1 page)", "Initial title proposal"],
     "Why this topic? Nepal healthcare context. Grande Hospital dataset plan. Supervisor meeting notes."),
    (2, "Literature Review & Requirements", "Phase 1: Research & Planning",
     ["Study CNN-based brain MRI classification papers",
      "Research explainable AI (Grad-CAM) for medical imaging",
      "Document functional requirements (upload, classify, report, roles)",
      "Document non-functional requirements (security, usability)",
      "Confirm Grande International Hospital, Kathmandu as MRI dataset source",
      "Create stakeholder list: radiologist, doctor, patient"],
     ["Literature review draft (5-10 sources)", "Requirements list", "Grande Hospital data access plan"],
     "Summary table of related work. Grande Hospital dataset governance notes."),
    (3, "System Design & Architecture", "Phase 1: Research & Planning",
     ["Draw system architecture (React, FastAPI, PyTorch pipeline)",
      "Design 8-stage pipeline flowchart",
      "Choose tech stack: Python, PyTorch, FastAPI, React, SQLite",
      "Plan folder structure and data flow",
      "Create Gantt chart for 32 weeks"],
     ["Architecture diagram", "Pipeline diagram", "Project plan"],
     "Design decisions and technology justification."),
    (4, "Ethics, Data & Project Proposal", "Phase 1: Research & Planning",
     ["Write ethics and data handling section",
      "Confirm Grande International Hospital anonymised MRI dataset as primary source",
      "Draft full project proposal",
      "Set up Git repository and .gitignore (exclude data/raw/ Grande Hospital scans)",
      "Prepare proposal presentation slides"],
     ["Project proposal document", "Ethics checklist", "Git repository"],
     "Grande Hospital anonymisation approach. Ethics approval. Repository URL."),
    (5, "Development Environment Setup", "Phase 2: Data & ML Foundation",
     ["Install Python 3.11, PyTorch, FastAPI, Node.js, React",
      "Create requirements.txt and requirements-cnn.txt",
      "Set up virtual environment",
      "Test GPU/CPU training capability",
      "Create project README with install instructions"],
     ["Working dev environment", "README"],
     "Screenshots of successful installs. Commands used. Issues faced."),
    (6, "Grande Hospital Dataset Acquisition & Exploration", "Phase 2: Data & ML Foundation",
     ["Receive Grande International Hospital anonymised MRI dataset (~1,600 scans)",
      "Organise into data/raw/normal and data/raw/abnormal",
      "Explore class distribution (Normal vs Abnormal)",
      "Visualise sample Grande Hospital scans",
      "Document image sizes, formats, quality issues",
      "Split data: train / validation / test"],
     ["Grande Hospital dataset report", "Sample images", "Split statistics"],
     "Grande Hospital dataset statistics table. Sample MRI figures with captions."),
    (7, "Preprocessing Pipeline", "Phase 2: Data & ML Foundation",
     ["Implement scan loading (JPEG, PNG, NIfTI)",
      "Implement CLAHE preprocessing",
      "Implement low-contrast quality check",
      "Resize/normalise to 128x128 grayscale",
      "Unit tests for preprocessing"],
     ["preprocessing.py", "Preprocessing test results"],
     "Before/after CLAHE images. QC threshold rationale."),
    (8, "Baseline CNN Model Design", "Phase 2: Data & ML Foundation",
     ["Design Baseline CNN architecture",
      "Implement forward pass and output layers",
      "Document model parameters",
      "Create model summary diagram"],
     ["Model architecture document", "cnn_baseline.py"],
     "Architecture diagram. Parameter count. Design choices."),
    (9, "Model Training (Phase 1)", "Phase 2: Data & ML Foundation",
     ["Implement training loop (train_cnn.py)",
      "Configure loss, optimiser, learning rate",
      "Run first training epochs; monitor loss",
      "Log training/validation accuracy",
      "Save checkpoints to models/cnn_baseline.pth"],
     ["Training logs", "Initial model weights"],
     "Training curves. Hardware used. Training time."),
    (10, "Model Evaluation & Tuning", "Phase 2: Data & ML Foundation",
     ["Evaluate on Grande Hospital validation/test sets (evaluate.py)",
      "Compute accuracy, precision, recall, F1 on Grande Hospital dataset",
      "Generate confusion matrix",
      "Tune hyperparameters if needed",
      "Document final validation accuracy (97.81% on Grande Hospital data)"],
     ["Evaluation report", "cnn_baseline_report.txt"],
     "Grande Hospital dataset metrics table. Confusion matrix. Comparison with literature."),
    (11, "Inference & Prediction Module", "Phase 3: AI Pipeline & Explainability",
     ["Implement predict.py for single MRI inference",
      "Load trained weights and run prediction",
      "Return label and confidence score",
      "Test on 10+ sample images",
      "Write test_predict.py"],
     ["Working predict script", "Test results"],
     "Sample predictions with confidence. Edge cases."),
    (12, "Grad-CAM Explainability", "Phase 3: AI Pipeline & Explainability",
     ["Implement Grad-CAM heatmap generation",
      "Overlay heatmap on original scan",
      "Save output to results/jobs/{id}/gradcam.png",
      "Test on normal and abnormal cases",
      "Document clinical trust rationale"],
     ["Grad-CAM module", "Sample heatmap images"],
     "Normal vs abnormal heatmap comparison."),
    (13, "RAG Advisory & Medical Knowledge Base", "Phase 3: AI Pipeline & Explainability",
     ["Build medical knowledge base (MEDICAL_KB)",
      "Implement RAG-style advisory for normal/abnormal",
      "Add disclaimer (prototype, not medical advice)",
      "Link advisory to confidence score",
      "Draft advisory text for both outcomes"],
     ["RAG advisory module", "Sample advisory outputs"],
     "Sample advisory text. Limitations of rule-based RAG."),
    (14, "Bilingual Chatbot & Hospital Finder", "Phase 3: AI Pipeline & Explainability",
     ["Implement English/Nepali chatbot responses",
      "Create Nepal hospital list (neurology centres)",
      "Filter hospitals by abnormal vs normal result",
      "Test bilingual output formatting"],
     ["Chatbot module", "Hospital finder module"],
     "Sample EN/NE responses. Hospital list source."),
    (15, "End-to-End Pipeline Integration", "Phase 3: AI Pipeline & Explainability",
     ["Implement pipeline.py with all 8 stages",
      "Add stage logging and callbacks",
      "Generate JSON/TXT/HTML reports",
      "Test full pipeline on 5 scans"],
     ["Complete pipeline.py", "Sample reports"],
     "Pipeline flow diagram. Stage timing. Sample report."),
    (16, "Pipeline Testing & Logging", "Phase 3: AI Pipeline & Explainability",
     ["Set up structured logging (neuroscan.log)",
      "Run integration test on 25 images",
      "Document pass/fail rate and processing time",
      "Fix pipeline bugs from integration test"],
     ["Integration test results", "Log file samples"],
     "Test summary table. Bugs found and fixed."),
    (17, "FastAPI Backend Foundation", "Phase 4: Backend API",
     ["Create backend.py with FastAPI",
      "Implement GET /health, GET /pipeline/steps",
      "Set up CORS for frontend",
      "Configure upload and results directories",
      "Test with Postman/curl"],
     ["Basic API", "Health endpoint working"],
     "API endpoint list. Postman screenshot."),
    (18, "Job Management & File Upload", "Phase 4: Backend API",
     ["Implement POST /upload",
      "Implement job store (jobs.json)",
      "Background thread for process_job()",
      "Implement GET /jobs, GET /jobs/{id}",
      "Store result.score, label, full pipeline output"],
     ["Upload API", "Job persistence"],
     "Job JSON schema. Upload flow diagram."),
    (19, "Backend Error Handling & Polish", "Phase 4: Backend API",
     ["Handle pipeline unavailable (missing PyTorch)",
      "Handle upload and job failures",
      "Add request logging middleware",
      "Document API in README or Swagger",
      "Test concurrent uploads"],
     ["Robust error handling", "API documentation"],
     "Error scenarios tested. Swagger screenshot."),
    (20, "Confidence Score Fix & Data Consistency", "Phase 4: Backend API",
     ["Fix confidence display (score vs confidence)",
      "Create jobResult.js utility",
      "Verify jobs.json matches UI display",
      "Regression test on existing jobs"],
     ["Fixed confidence display"],
     "Bug description, root cause, fix applied."),
    (21, "React App Setup & Routing", "Phase 5: Frontend Development",
     ["Set up React + Vite + Tailwind",
      "Create routes: Home, Upload, Processing, Results, Model, About",
      "Build Header, Footer, PageHeader, StatusBadge",
      "Configure API_BASE and pipeline steps"],
     ["React app skeleton", "Navigation working"],
     "UI wireframe vs implementation."),
    (22, "Upload & Processing Pages", "Phase 5: Frontend Development",
     ["Build Upload page (file picker, preview, progress)",
      "Implement XHR upload with progress",
      "Build Processing page with 8-stage live status",
      "Poll GET /jobs/{id} every 2 seconds",
      "Show pipeline progress bar"],
     ["Upload + Processing pages functional"],
     "Upload UX flow. Polling strategy. Screenshots."),
    (23, "Results & Dashboard Pages", "Phase 5: Frontend Development",
     ["Build Results page (classification, RAG, hospitals)",
      "Build Home dashboard (job stats, recent activity)",
      "Build Model Info and About pages",
      "Backend status indicator component"],
     ["Results page", "Home dashboard"],
     "Results page layout. Sample completed job screenshot."),
    (24, "UI/UX Refinement", "Phase 5: Frontend Development",
     ["Apply consistent styling (cards, buttons, gradients)",
      "Responsive design for mobile/tablet",
      "Status badges for job states",
      "Improve error messages and empty states",
      "User testing with peers; collect feedback"],
     ["Polished UI", "Feedback notes"],
     "Before/after UI. Peer feedback summary."),
    (25, "Prototype HTML Page", "Phase 5: Frontend Development",
     ["Build/update prototype.html for quick demos",
      "Sample cases (normal/abnormal) without upload",
      "Live upload with backend integration",
      "Carousel for sample MRI images"],
     ["Working prototype page"],
     "Demo script for presentations."),
    (26, "Frontend-Backend Integration Testing", "Phase 5: Frontend Development",
     ["Full flow: upload -> processing -> results",
      "Test multiple file formats (JPEG, PNG)",
      "Verify report download paths",
      "Fix CORS/auth issues",
      "Document known limitations"],
     ["E2E test checklist", "Bug fixes"],
     "Test cases passed/failed. Integration issues resolved."),
    (27, "Authentication & User Roles", "Phase 6: RBAC & Clinical Workflow",
     ["Implement src/auth.py (SQLite, bcrypt, JWT)",
      "Roles: radiologist, doctor, patient",
      "Auto-generate patient unique ID (NSN-PAT-000001)",
      "Login/Register pages, AuthContext, ProtectedRoute",
      "Seed demo accounts"],
     ["Auth system", "Login/register UI"],
     "RBAC matrix. Security measures."),
    (28, "Role-Based Workflows", "Phase 6: RBAC & Clinical Workflow",
     ["Radiologist: upload with patient assignment",
      "Doctor: MRI Review page (scan, AI, recommendation)",
      "Patient: My Records portal",
      "API: doctor-review, scan-image endpoints",
      "Filter jobs by role on backend"],
     ["Doctor Review", "Patient Portal", "Role-filtered API"],
     "Workflow diagram per role. Screenshots."),
    (29, "Radiologist Notes & Prototype Auth Fix", "Phase 6: RBAC & Clinical Workflow",
     ["Add radiologist investigation recommendation",
      "POST /jobs/{id}/radiologist-notes",
      "Show radiologist notes to doctor and patient",
      "Fix prototype.html 401 auth issue",
      "Full demo: register -> upload -> notes -> doctor -> patient"],
     ["Radiologist notes feature", "Fixed prototype"],
     "Full clinical workflow demo script."),
    (30, "System Testing & Performance", "Phase 7: Testing & Submission",
     ["Run integration test on 25+ scans",
      "Test all three roles (login, permissions, isolation)",
      "Measure average pipeline time per scan",
      "Test backend offline graceful degradation",
      "Create test report document"],
     ["Test report", "Performance metrics"],
     "Test summary. Performance table."),
    (31, "Documentation & Deployment Guide", "Phase 7: Testing & Submission",
     ["Complete README (install, run, demo accounts)",
      "User manual for radiologist, doctor, patient",
      "Technical report sections: design, implementation, evaluation",
      "Architecture and pipeline diagrams",
      "Prepare demo video or live demo script"],
     ["README", "User manual", "Report drafts"],
     "Documentation checklist. Demo rehearsal notes."),
    (32, "Final Review & Submission", "Phase 7: Testing & Submission",
     ["Final code cleanup and Git commit",
      "Proofread report and logbook",
      "Prepare viva/presentation slides",
      "Rehearse demo with all three roles",
      "Submit report, logbook, code, and forms",
      "Reflect on outcomes and future work"],
     ["Final submission package", "Presentation"],
     "Final reflection. Achievements, challenges, future improvements."),
]


def set_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    return h


def add_bullet_list(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_logbook_template(doc, week_num, title):
    doc.add_paragraph()
    table = doc.add_table(rows=8, cols=2)
    table.style = "Table Grid"
    rows = [
        ("Week / Title", f"Week {week_num}: {title}"),
        ("Date Range", "___________________________"),
        ("Hours Spent", "___________________________"),
        ("Objectives", ""),
        ("Activities Completed", ""),
        ("Challenges & Solutions", ""),
        ("Skills Learned", ""),
        ("Supervisor Comments", ""),
    ]
    for i, (label, value) in enumerate(rows):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
        for cell in table.rows[i].cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
    doc.add_paragraph()


def build_document():
    doc = Document()

    # Title page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("NeuroScan Nepal\n")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("Final Year Project — 32-Week Logbook Plan\n")
    r2.font.size = Pt(16)
    r2.bold = True

    sub2 = doc.add_paragraph()
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = sub2.add_run("Brain MRI Screening & Clinical Support System\n")
    r3.font.size = Pt(14)

    doc.add_paragraph()
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in [
        "Student Name: _________________________________",
        "Student ID: ___________________________________",
        "Supervisor: ___________________________________",
        "Institution: __________________________________",
        "Academic Year: ________________________________",
    ]:
        info.add_run(line + "\n").font.size = Pt(11)

    doc.add_page_break()

    # Table of contents style overview
    set_heading(doc, "Project Overview", 1)
    doc.add_paragraph(
        "NeuroScan Nepal is an AI-assisted brain MRI screening system designed for Nepal's "
        "healthcare context. It combines deep learning (CNN), explainable AI (Grad-CAM), "
        "clinical advisory (RAG), and role-based access for radiologists, doctors, and patients."
    )

    set_heading(doc, "Primary Dataset", 2)
    doc.add_paragraph(
        "Grande International Hospital, Kathmandu — approximately 1,600 anonymised brain MRI scans "
        "organised into normal and abnormal classes. Stored locally in data/raw/normal and "
        "data/raw/abnormal. All patient identifiers removed per BCU ethics approval. The dataset "
        "is not committed to Git. The hospital finder module also recommends Grande International "
        "Hospital for neurology referrals in Kathmandu."
    )

    set_heading(doc, "Technology Stack", 2)
    add_bullet_list(doc, [
        "Backend: Python 3.11, FastAPI, PyTorch, SQLite",
        "Frontend: React 18, Vite, Tailwind CSS",
        "ML: Baseline CNN, CLAHE preprocessing, Grad-CAM",
        "Auth: JWT, bcrypt, role-based access control",
        "Storage: jobs.json, uploads/, results/jobs/",
    ])

    set_heading(doc, "Project Milestones", 2)
    milestones = doc.add_table(rows=9, cols=2)
    milestones.style = "Table Grid"
    ms_data = [
        ("Milestone", "Target Week"),
        ("Proposal approved", "Week 4"),
        ("Dataset & preprocessing ready", "Week 7"),
        ("Model trained & evaluated", "Week 10"),
        ("Full AI pipeline working", "Week 16"),
        ("Backend API complete", "Week 20"),
        ("Frontend complete", "Week 26"),
        ("RBAC & clinical workflow complete", "Week 29"),
        ("Final submission", "Week 32"),
    ]
    for i, (a, b) in enumerate(ms_data):
        milestones.rows[i].cells[0].text = a
        milestones.rows[i].cells[1].text = b
        if i == 0:
            for cell in milestones.rows[i].cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True

    set_heading(doc, "Demo Accounts (Week 27+)", 2)
    add_bullet_list(doc, [
        "Radiologist: radiologist@neuroscan.np / radiologist123",
        "Doctor: doctor@neuroscan.np / doctor123",
        "Patient: patient@neuroscan.np / patient123",
    ])

    doc.add_page_break()

    # Phase summary
    set_heading(doc, "Phase Summary", 1)
    phases = [
        ("Weeks 1–4", "Project Initiation & Research"),
        ("Weeks 5–10", "Data & ML Foundation"),
        ("Weeks 11–16", "AI Pipeline & Explainability"),
        ("Weeks 17–20", "Backend API Development"),
        ("Weeks 21–26", "Frontend Development"),
        ("Weeks 27–29", "RBAC & Clinical Workflow"),
        ("Weeks 30–32", "Testing, Documentation & Submission"),
    ]
    for weeks, phase in phases:
        p = doc.add_paragraph()
        p.add_run(f"{weeks}: ").bold = True
        p.add_run(phase)

    doc.add_page_break()

    # Weekly breakdown
    set_heading(doc, "Weekly Breakdown & Logbook Entries", 1)

    current_phase = None
    for week_num, title, phase, tasks, deliverables, logbook_note in WEEKS:
        if phase != current_phase:
            current_phase = phase
            doc.add_page_break()
            set_heading(doc, phase, 1)

        set_heading(doc, f"Week {week_num}: {title}", 2)

        set_heading(doc, "Tasks / Activities", 3)
        add_bullet_list(doc, tasks)

        set_heading(doc, "Deliverables", 3)
        add_bullet_list(doc, deliverables)

        set_heading(doc, "Logbook Notes (What to Write)", 3)
        doc.add_paragraph(logbook_note)

        set_heading(doc, "Weekly Logbook Entry Template", 3)
        add_logbook_template(doc, week_num, title)

        if week_num < 32:
            doc.add_paragraph("—" * 60)

    doc.add_page_break()

    # Appendix
    set_heading(doc, "Appendix A: Logbook Writing Tips", 1)
    tips = doc.add_table(rows=8, cols=2)
    tips.style = "Table Grid"
    tip_data = [
        ("Element", "What to Include"),
        ("Date & Hours", "Date range and hours spent (8-12 hrs/week typical)"),
        ("Activities", "Bullet list of tasks completed"),
        ("Evidence", "Screenshots, diagrams, code snippets, test results"),
        ("Problems", "Issues faced and how you solved them"),
        ("Learning", "New skills acquired (PyTorch, FastAPI, RBAC, etc.)"),
        ("Next Week", "Planned tasks for following week"),
        ("Supervisor", "Meeting date and feedback received"),
    ]
    for i, (a, b) in enumerate(tip_data):
        tips.rows[i].cells[0].text = a
        tips.rows[i].cells[1].text = b

    set_heading(doc, "Appendix B: Clinical Workflow (Final System)", 1)
    doc.add_paragraph(
        "1. Patient registers → receives unique ID (NSN-PAT-000001)\n"
        "2. Radiologist uploads MRI → assigns to patient → AI pipeline runs\n"
        "3. Radiologist adds investigation recommendation on Results page\n"
        "4. Doctor reviews MRI on MRI Review page → submits recommendation\n"
        "5. Patient views MRI, AI report, and doctor recommendation on My Records"
    )

    set_heading(doc, "Appendix C: 8-Stage AI Pipeline", 1)
    add_bullet_list(doc, [
        "1. Upload — Save MRI scan to server",
        "2. Preprocessing — CLAHE + quality check",
        "3. Detection — CNN classification (Normal/Abnormal)",
        "4. Grad-CAM — Explainability heatmap",
        "5. RAG Advisory — Medical knowledge advisory text",
        "6. Chatbot — Bilingual patient guidance (EN/NE)",
        "7. Hospital Finder — Nepal neurology centre recommendations",
        "8. Report Generation — JSON, TXT, HTML export",
    ])

    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    build_document()
